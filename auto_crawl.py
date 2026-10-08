#!/usr/bin/python
# coding:utf-8 -*-
"""抖音直播自动全程采集（守护模式：弹幕/事件/商品讲解，可选音视频）

从 xlsx 读取主播列表（抖音号），长期守护运行：
- 轮询检测每个主播是否开播（解析 room_id + 直播间状态）
- 检测到开播 → 自动启动 WebSocket 采集（弹幕/事件/商品 promotion_id）
- 主播下播（ControlMessage status==3）→ 自动停止并落盘（room.json/events.jsonl/...）
- 继续盯守，等主播下一次开播再采，循环往复直到手动停止

**音视频默认关闭**：`--media` 打开后每场额外起 ffmpeg 抽帧 + 切片（走
`liveMan.MediaCapture`，与 `collect_full.py` 同源）。默认关是因为音视频的
CPU/内存/磁盘开销**无法线性外推到 100 并发**（100 路软解约 50-100 核、
磁盘 20-30 GB/小时），只适合 5-10 个主播的小规模全模态验证。

用法：
  python auto_crawl.py [xlsx_path] [check_interval_seconds] [--sheet N] [--media]
  python auto_crawl.py documents/并发测试.xlsx 60 --sheet 40    # 取 40 并发档名单
  python auto_crawl.py anchor.xlsx 60 --media                  # 3-5 个主播带音视频
"""
import argparse
import io
import random
import sys
import threading
import time
from datetime import datetime

from liveMan import DouyinLiveWebFetcher, MediaCapture
from data_recorder import DataRecorder, new_date_dir

CHECK_INTERVAL = 60  # 默认开播检测间隔（秒）
MEDIA_MAX_SECONDS = 8 * 3600  # 音视频的防僵尸上限（真正的停止是 stop_media()）


# 表头候选名：不同名单文件措辞不同，按名字定位列而不是写死列号。
# 仓库里有两份结构不同的名单：
#   anchor.xlsx             10 列：分层 | 博主 | 9月粉丝 | … | 抖音号 | 直播间链接
#   documents/并发测试.xlsx   5 列：序号 | 昵称 | 抖音号 | 粉丝数 | 主要类目
# 旧实现硬编码 `r[8]`（第 9 列）取抖音号，对后者会**静默返回 0 个主播**（5 列文件里
# r[8] 恒为 None），排查起来很费时间 —— 故改为表头驱动，找不到列就明确报错。
ID_HEADER_CANDIDATES = ('抖音号', '抖音id', '抖音 ID', '账号', 'live_id')
NAME_HEADER_CANDIDATES = ('昵称', '博主', '博主名', '主播', '主播名', '名字', 'name')


def _norm_cell(v):
    """单元格值归一化为去空白的字符串（None → ''）"""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _find_col(header, candidates):
    """在表头里找候选列名，返回列索引；找不到返回 None（做子串匹配以容忍空格/前缀）"""
    for i, h in enumerate(header):
        hs = _norm_cell(h).lower()
        if not hs:
            continue
        for cand in candidates:
            c = cand.lower()
            if hs == c or c in hs:
                return i
    return None


def load_anchors(xlsx_path, sheet=None):
    """读取主播名单，返回 [(name, douyin_id)]

    **按表头名定位列**，不写死列号（见 ID_HEADER_CANDIDATES 的说明）。

    :param sheet: sheet 名，如 `'40'` 取 `并发测试.xlsx` 的 40 并发档名单；
                  默认取第一个 sheet。分档 sheet 与 sheet1 列布局相同但没有表头行，
                  此时用 `sheet1` 解析列位置，并自动跳过目标 sheet 的首行（若它是表头）。
    :raises ValueError: 找不到抖音号列或昵称列（明确报错，不静默返回空列表）
    """
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    # 表头来源：优先 sheet1（并发测试.xlsx 的分档表没有表头），否则用目标/活动 sheet
    header_ws = wb['sheet1'] if 'sheet1' in wb.sheetnames else (
        wb[str(sheet)] if sheet is not None and str(sheet) in wb.sheetnames else wb.active)
    header = [c for c in next(header_ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    id_idx = _find_col(header, ID_HEADER_CANDIDATES)
    name_idx = _find_col(header, NAME_HEADER_CANDIDATES)
    if id_idx is None:
        raise ValueError(f"{xlsx_path} 的表头 {header} 里找不到抖音号列"
                         f"（候选：{ID_HEADER_CANDIDATES}）")
    if name_idx is None:
        raise ValueError(f"{xlsx_path} 的表头 {header} 里找不到昵称/博主列"
                         f"（候选：{NAME_HEADER_CANDIDATES}）")

    ws = wb[str(sheet)] if sheet is not None else wb.active
    rows = list(ws.iter_rows(values_only=True))
    # 首行是否表头：抖音号列的值命中候选名就是表头（分档 sheet 首行即数据，需保留）
    start = 1 if rows and _norm_cell(rows[0][id_idx]).lower() in \
        [c.lower() for c in ID_HEADER_CANDIDATES] else 0

    anchors = []
    for r in rows[start:]:
        if not r or all(c is None or str(c).strip() == '' for c in r):
            continue
        name = _norm_cell(r[name_idx]) if name_idx < len(r) else ''
        douyin_id = _norm_cell(r[id_idx]) if id_idx < len(r) else ''
        # 抖音号可能是纯数字（int/float）或字母数字串，_norm_cell 已统一处理
        if name and douyin_id and douyin_id.lower() != 'none':
            anchors.append((name, douyin_id))
    return anchors


class AnchorDaemon(threading.Thread):
    """单个主播的守护线程：检测开播 → 全程采集 → 落盘 → 循环"""

    def __init__(self, name, douyin_id, check_interval=CHECK_INTERVAL, media=False):
        super().__init__(name=name)  # 非 daemon：保证优雅退出时 close 完成
        self.anchor_name = name
        self.douyin_id = douyin_id
        self.check_interval = check_interval
        self.media = media          # 是否采音视频（默认关，见模块 docstring）
        self._stop_event = threading.Event()

    def stop(self):
        self._stop_event.set()

    def run(self):
        # 随机初始延迟：避免多个线程同时首次请求，触发抖音风控（HTTP 444）
        init_delay = random.uniform(0, self.check_interval)
        print(f"[{self.anchor_name}] 守护启动（抖音号 {self.douyin_id}，检测间隔 {self.check_interval}s，初始延迟 {init_delay:.1f}s）")
        self._stop_event.wait(init_delay)
        while not self._stop_event.is_set():
            try:
                self._one_cycle()
            except Exception as e:
                print(f"[{self.anchor_name}] 循环异常: {e}")
            # 加随机抖动，避免所有线程同步轮询造成突发并发
            jitter = random.uniform(0.5, 1.5)
            self._stop_event.wait(self.check_interval * jitter)

    def _one_cycle(self):
        """一轮检测：若开播则采集到自然下播"""
        fetcher = DouyinLiveWebFetcher(self.douyin_id, None)

        # 1. 解析 room_id（抖音号有效即可解析；未开播时返回历史房间号，开播判断靠下一步）
        try:
            rid = fetcher.room_id
        except Exception:
            rid = None
        if not rid:
            return  # 抖音号无效，跳过本轮

        # 2. 确认直播状态 + 主播信息
        try:
            success, status, nickname, user_id = fetcher.get_room_status()
        except Exception:
            return
        if not success or status != '正在直播':
            return

        start_ts = datetime.now()
        print(f"[{self.anchor_name}] {start_ts.strftime('%H:%M:%S')} 检测到开播：{nickname}（{user_id}），开始全程采集")

        # 3. 启动采集（WebSocket 弹幕 + 结构化落盘）
        offline = threading.Event()
        fetcher.on_offline = offline.set

        # 落盘目录：`data/<当天日期>/`（见 data_recorder.new_date_dir）。在**确认开播之后**
        # 才解析 —— 若在检测前就解析，一次没人开播的轮询也会建出一个空目录，反而更乱。
        # 日期在这一刻定死：本场跨零点也继续写同一目录，直到下播，不清场不换文件夹。
        recorder = DataRecorder(fetcher.room_id, self.douyin_id, base_dir=new_date_dir())
        room_info = dict(fetcher.room_detail)  # enter API 富字段（头像/简介/粉丝数/作品数等）
        room_info.update({
            'title': nickname, 'anchor_nickname': nickname,
            'anchor_id': user_id, 'room_status': status,
        })
        recorder.set_room_info(room_info)
        fetcher.data_callback = recorder.on_event

        ws_thread = threading.Thread(target=fetcher.start, daemon=True)
        ws_thread.start()

        # 3.5 音视频（默认关）：与 collect_full 共用 MediaCapture，
        # 区别只在**谁调 stop_media()** —— 这里挂到「下播」这一自然终点。
        capture = None
        if self.media:
            capture = MediaCapture(self.douyin_id, recorder.path, fetcher=fetcher,
                                   max_seconds=MEDIA_MAX_SECONDS)
            if not capture.start_media():
                print(f"[{self.anchor_name}] ⚠️ 音视频启动失败，本场只采文本/数值")

        # 4. 等待下播（ControlMessage status==3 触发 on_offline）
        while not offline.is_set() and not self._stop_event.is_set():
            if not ws_thread.is_alive():
                break  # ws 意外断开（网络问题），也结束本场
            offline.wait(timeout=self.check_interval)

        # 5. 停止并落盘（先停音视频再关 recorder，避免 ffmpeg 还在往目录里写）
        media = None
        if capture is not None:
            media = capture.stop_media()
        fetcher.stop()
        ws_thread.join(timeout=5)  # 等 WS 线程退出，否则 close 后仍可能有回调写入
        recorder.set_promotion_order(fetcher.product_refresh_order)
        recorder.close()
        if media:
            print(f"[{self.anchor_name}] 音视频：{len(media['frames'])} 帧 / "
                  f"{len(media['audio'])} 段音频"
                  + (f"，续采 {media['restarts']} 次" if media['restarts'] else ""))
        dur = (datetime.now() - start_ts).total_seconds()
        print(f"[{self.anchor_name}] 本场结束（时长 {dur / 60:.1f} 分钟），数据落盘：{recorder.path}")


def main():
    try:
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace', line_buffering=True)
    except Exception:
        pass

    ap = argparse.ArgumentParser(
        description="抖音直播守护式采集（多主播轮播长跑）",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    ap.add_argument('xlsx_path', nargs='?', default='anchor.xlsx',
                    help='主播名单 xlsx（表头驱动定位「抖音号」/「昵称」列）')
    ap.add_argument('check_interval', nargs='?', type=int, default=CHECK_INTERVAL,
                    help='开播检测间隔（秒）')
    ap.add_argument('--sheet', default=None,
                    help='取该 sheet 的名单，如 40 取 并发测试.xlsx 的 40 并发档；'
                         '默认取第一个 sheet')
    ap.add_argument('--media', action='store_true',
                    help='同时采音视频（默认关；开销无法线性外推到 100 并发，'
                         '只适合 5-10 个主播）')
    args = ap.parse_args()

    anchors = load_anchors(args.xlsx_path, sheet=args.sheet)
    if not anchors:
        print(f"错误：未从 {args.xlsx_path} 读取到主播列表")
        return
    print(f"加载主播列表 {len(anchors)} 个（来自 {args.xlsx_path}"
          f"{'，sheet ' + str(args.sheet) if args.sheet else ''}，"
          f"检测间隔 {args.check_interval}s，音视频 {'开' if args.media else '关'}）：")
    for name, did in anchors:
        print(f"  - {name}（抖音号 {did}）")
    print("=" * 70)

    # 无需共享的目录对象：每个主播在自己开播时各自解析 `data/<当天日期>/`，
    # 同名目录由 os.makedirs(exist_ok=True) 幂等收敛（见 data_recorder.new_date_dir）
    daemons = [AnchorDaemon(name, did, args.check_interval, media=args.media)
               for name, did in anchors]
    for d in daemons:
        d.start()

    print("守护模式运行中，按 Ctrl+C 停止...")
    try:
        while True:
            time.sleep(3600)
    except KeyboardInterrupt:
        print("\n收到停止信号，正在优雅关闭...")
        for d in daemons:
            d.stop()
        for d in daemons:
            d.join(timeout=30)
        print("已全部停止，退出。")


if __name__ == "__main__":
    main()
